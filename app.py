from flask import Flask, request, jsonify, render_template, redirect, url_for
from datetime import datetime, date
import sqlite3, hashlib, math, os, pytz

app = Flask(__name__)

TURKEY_TZ        = pytz.timezone("Europe/Istanbul")
DB_PATH          = os.environ.get("DB_PATH", "yoklama.db")
ADMIN_PASS       = os.environ.get("ADMIN_PASSWORD", "admin123")
KLINIK_LAT       = float(os.environ.get("KLINIK_LAT",       "40.9838647"))
KLINIK_LNG       = float(os.environ.get("KLINIK_LNG",       "27.5695521"))
KLINIK_YARICAP_M = int(os.environ.get("KLINIK_YARICAP_M",   "100"))

# ── MESAI TAKVİMİ ─────────────────────────────────────────────────────────────
# weekday: 0=Pzt … 4=Cuma, 5=Cmt, 6=Pzr
MESAI = {
    "hafta_ici": {"giris": (9, 0), "cikis": (18, 0)},
    "cumartesi": {"giris": (9, 0), "cikis": (14, 0)},
}
TOLERANS_DK = 60

# ── VERİTABANI ────────────────────────────────────────────────────────────────
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS personel (
                id               INTEGER PRIMARY KEY AUTOINCREMENT,
                ad_soyad         TEXT    NOT NULL,
                pin_hash         TEXT    NOT NULL,
                aktif            INTEGER DEFAULT 1,
                olusturma_tarihi TEXT    DEFAULT (datetime('now'))
            );
            CREATE TABLE IF NOT EXISTS yoklama (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                personel_id   INTEGER NOT NULL,
                tarih         TEXT    NOT NULL,
                giris_saati   TEXT,
                cikis_saati   TEXT,
                giris_ip      TEXT,
                cikis_ip      TEXT,
                giris_lat     REAL,
                giris_lng     REAL,
                cikis_lat     REAL,
                cikis_lng     REAL,
                giris_uyari   INTEGER DEFAULT 0,
                cikis_uyari   INTEGER DEFAULT 0,
                cikis_disari  INTEGER DEFAULT 0,
                FOREIGN KEY (personel_id) REFERENCES personel(id)
            );
        """)
        # Otomatik migration — eski DB'lere eksik kolon ekle
        mevcut = [r[1] for r in conn.execute("PRAGMA table_info(yoklama)").fetchall()]
        ekstra = {
            "giris_lat":    "ALTER TABLE yoklama ADD COLUMN giris_lat    REAL",
            "giris_lng":    "ALTER TABLE yoklama ADD COLUMN giris_lng    REAL",
            "cikis_lat":    "ALTER TABLE yoklama ADD COLUMN cikis_lat    REAL",
            "cikis_lng":    "ALTER TABLE yoklama ADD COLUMN cikis_lng    REAL",
            "giris_uyari":  "ALTER TABLE yoklama ADD COLUMN giris_uyari  INTEGER DEFAULT 0",
            "cikis_uyari":  "ALTER TABLE yoklama ADD COLUMN cikis_uyari  INTEGER DEFAULT 0",
            "cikis_disari": "ALTER TABLE yoklama ADD COLUMN cikis_disari INTEGER DEFAULT 0",
        }
        for kolon, sql in ekstra.items():
            if kolon not in mevcut:
                conn.execute(sql)
        conn.commit()

# ── YARDIMCI FONKSİYONLAR ────────────────────────────────────────────────────
def now_tr():
    return datetime.now(TURKEY_TZ)

def hash_pin(pin):
    return hashlib.sha256(pin.strip().encode()).hexdigest()

def get_personel(pin):
    with get_db() as conn:
        return conn.execute(
            "SELECT * FROM personel WHERE pin_hash=? AND aktif=1",
            (hash_pin(pin),)
        ).fetchone()

def bugun_yoklama(personel_id, tarih):
    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM yoklama WHERE personel_id=? AND tarih=?",
            (personel_id, tarih)
        ).fetchone()
        if row:
            return dict(row)
        conn.execute(
            "INSERT INTO yoklama (personel_id, tarih) VALUES (?,?)",
            (personel_id, tarih)
        )
        conn.commit()
        return dict(conn.execute(
            "SELECT * FROM yoklama WHERE personel_id=? AND tarih=?",
            (personel_id, tarih)
        ).fetchone())

def haversine(lat1, lng1, lat2, lng2):
    R, p = 6_371_000, math.pi / 180
    a = (math.sin((lat2-lat1)*p/2)**2
         + math.cos(lat1*p)*math.cos(lat2*p)*math.sin((lng2-lng1)*p/2)**2)
    return round(2 * R * math.asin(math.sqrt(a)))

def konum_kontrol(lat, lng):
    mesafe = haversine(KLINIK_LAT, KLINIK_LNG, lat, lng)
    return mesafe <= KLINIK_YARICAP_M, mesafe

def mesai_kontrol(now, islem):
    gun = now.weekday()
    dk  = now.hour * 60 + now.minute
    if gun == 6:
        return True, "⚠️ Bugün Pazar — normal mesai günü değil."
    mesai = MESAI["cumartesi"] if gun == 5 else MESAI["hafta_ici"]
    if islem == "giris":
        sinir = mesai["giris"][0]*60 + mesai["giris"][1] - TOLERANS_DK
        if dk < sinir:
            return True, f"⚠️ Mesai başlangıcından {TOLERANS_DK} dk'dan fazla önce giriş. Normal giriş: {mesai['giris'][0]:02d}:{mesai['giris'][1]:02d}"
    else:
        sinir = mesai["cikis"][0]*60 + mesai["cikis"][1] + TOLERANS_DK
        if dk > sinir:
            return True, f"⚠️ Mesai bitişinden {TOLERANS_DK} dk'dan fazla sonra çıkış. Normal çıkış: {mesai['cikis'][0]:02d}:{mesai['cikis'][1]:02d}"
    return False, None

def admin_kontrol(sifre):
    return sifre == ADMIN_PASS

def sure_hesapla(giris, cikis, tarih):
    try:
        g = datetime.strptime(f"{tarih} {giris}", "%Y-%m-%d %H:%M:%S")
        c = datetime.strptime(f"{tarih} {cikis}", "%Y-%m-%d %H:%M:%S")
        dk = int((c - g).total_seconds() / 60)
        return dk, f"{dk//60}s {dk%60}dk"
    except:
        return 0, "-"

# ── PERSONEL ROUTES ───────────────────────────────────────────────────────────
@app.route("/giris")
def giris_page():
    return render_template("islem.html", islem="giris")

@app.route("/cikis")
def cikis_page():
    return render_template("islem.html", islem="cikis")

@app.route("/api/giris", methods=["POST"])
def api_giris():
    data = request.get_json()
    pin  = data.get("pin", "").strip()
    lat  = data.get("lat")
    lng  = data.get("lng")

    if not pin:
        return jsonify({"ok": False, "mesaj": "PIN boş olamaz."})

    # Konum zorunlu
    if lat is None or lng is None:
        return jsonify({"ok": False, "mesaj": "📍 Konum alınamadı. Lütfen konum iznini verin.", "konum_hatasi": True})
    try:
        lat, lng = float(lat), float(lng)
    except:
        return jsonify({"ok": False, "mesaj": "Geçersiz konum.", "konum_hatasi": True})

    # Giriş için klinik içi zorunlu
    iceride, mesafe = konum_kontrol(lat, lng)
    if not iceride:
        return jsonify({"ok": False, "mesaj": f"🚫 Kliniğe çok uzaktasınız ({mesafe} m). Giriş yalnızca klinik içinden yapılabilir.", "konum_hatasi": True})

    personel = get_personel(pin)
    if not personel:
        return jsonify({"ok": False, "mesaj": "Hatalı PIN."})

    now   = now_tr()
    tarih = now.strftime("%Y-%m-%d")
    saat  = now.strftime("%H:%M:%S")
    yoklama = bugun_yoklama(personel["id"], tarih)

    if yoklama["giris_saati"]:
        return jsonify({"ok": False, "mesaj": f"Bugün zaten giriş yaptınız ({yoklama['giris_saati'][:5]})."})

    uyari, uyari_mesaj = mesai_kontrol(now, "giris")
    ip = request.headers.get("X-Forwarded-For", request.remote_addr)

    with get_db() as conn:
        conn.execute(
            "UPDATE yoklama SET giris_saati=?, giris_ip=?, giris_lat=?, giris_lng=?, giris_uyari=? WHERE id=?",
            (saat, ip, lat, lng, 1 if uyari else 0, yoklama["id"])
        )
        conn.commit()

    return jsonify({
        "ok": True,
        "mesaj": f"Günaydın, {personel['ad_soyad']}! Giriş saatiniz: {saat[:5]}",
        "uyari": uyari,
        "uyari_mesaj": uyari_mesaj
    })

@app.route("/api/cikis", methods=["POST"])
def api_cikis():
    data = request.get_json()
    pin  = data.get("pin", "").strip()
    lat  = data.get("lat")
    lng  = data.get("lng")

    if not pin:
        return jsonify({"ok": False, "mesaj": "PIN boş olamaz."})

    if lat is None or lng is None:
        return jsonify({"ok": False, "mesaj": "📍 Konum alınamadı. Lütfen konum iznini verin.", "konum_hatasi": True})
    try:
        lat, lng = float(lat), float(lng)
    except:
        return jsonify({"ok": False, "mesaj": "Geçersiz konum.", "konum_hatasi": True})

    # Çıkış için konum kontrolü — dışarıdaysa uyar ama engelleme
    iceride, mesafe = konum_kontrol(lat, lng)
    cikis_disari = 0 if iceride else 1

    personel = get_personel(pin)
    if not personel:
        return jsonify({"ok": False, "mesaj": "Hatalı PIN."})

    now   = now_tr()
    tarih = now.strftime("%Y-%m-%d")
    saat  = now.strftime("%H:%M:%S")
    yoklama = bugun_yoklama(personel["id"], tarih)

    if not yoklama["giris_saati"]:
        return jsonify({"ok": False, "mesaj": "Bugün giriş kaydınız bulunamadı."})
    if yoklama["cikis_saati"]:
        return jsonify({"ok": False, "mesaj": f"Bugün zaten çıkış yaptınız ({yoklama['cikis_saati'][:5]})."})

    uyari, uyari_mesaj = mesai_kontrol(now, "cikis")

    # Klinik dışı çıkış uyarısı
    if cikis_disari:
        uyari = True
        disari_mesaj = f"⚠️ Klinik dışında çıkış yapıldı ({mesafe} m uzakta). Kayıt 'ihlal' olarak işaretlendi."
        uyari_mesaj  = disari_mesaj + (" | " + uyari_mesaj if uyari_mesaj else "")

    ip = request.headers.get("X-Forwarded-For", request.remote_addr)
    with get_db() as conn:
        conn.execute(
            "UPDATE yoklama SET cikis_saati=?, cikis_ip=?, cikis_lat=?, cikis_lng=?, cikis_uyari=?, cikis_disari=? WHERE id=?",
            (saat, ip, lat, lng, 1 if uyari else 0, cikis_disari, yoklama["id"])
        )
        conn.commit()

    _, sure_text = sure_hesapla(yoklama["giris_saati"], saat, tarih)
    return jsonify({
        "ok": True,
        "mesaj": f"İyi günler, {personel['ad_soyad']}! Çıkış: {saat[:5]} | Süre: {sure_text}",
        "uyari": uyari,
        "uyari_mesaj": uyari_mesaj
    })

# ── ADMIN ROUTES ──────────────────────────────────────────────────────────────
@app.route("/admin")
def admin_giris():
    return render_template("admin_giris.html")

@app.route("/admin/<sifre>")
def admin_panel(sifre):
    if not admin_kontrol(sifre):
        return "<h2 style='font-family:sans-serif;color:red;padding:40px'>❌ Erişim Reddedildi</h2>", 403
    return render_template("admin_panel.html", s=sifre)

@app.route("/api/admin/personeller")
def api_personeller():
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    with get_db() as conn:
        rows = conn.execute(
            "SELECT id, ad_soyad, aktif, olusturma_tarihi FROM personel ORDER BY ad_soyad"
        ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/admin/personel_ekle", methods=["POST"])
def api_personel_ekle():
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    data = request.get_json()
    ad, pin = data.get("ad_soyad","").strip(), data.get("pin","").strip()
    if not ad or not pin:
        return jsonify({"ok": False, "mesaj": "Ad ve PIN zorunludur."})
    if len(pin) < 4:
        return jsonify({"ok": False, "mesaj": "PIN en az 4 haneli olmalı."})
    with get_db() as conn:
        conn.execute("INSERT INTO personel (ad_soyad, pin_hash) VALUES (?,?)", (ad, hash_pin(pin)))
        conn.commit()
    return jsonify({"ok": True, "mesaj": f"{ad} eklendi."})

@app.route("/api/admin/personel_sil/<int:pid>", methods=["POST"])
def api_personel_sil(pid):
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    with get_db() as conn:
        conn.execute("UPDATE personel SET aktif=0 WHERE id=?", (pid,))
        conn.commit()
    return jsonify({"ok": True})

@app.route("/api/admin/pin_degistir/<int:pid>", methods=["POST"])
def api_pin_degistir(pid):
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    yeni = request.get_json().get("pin","").strip()
    if len(yeni) < 4:
        return jsonify({"ok": False, "mesaj": "PIN en az 4 haneli olmalı."})
    with get_db() as conn:
        conn.execute("UPDATE personel SET pin_hash=? WHERE id=?", (hash_pin(yeni), pid))
        conn.commit()
    return jsonify({"ok": True, "mesaj": "PIN güncellendi."})

@app.route("/api/admin/yoklamalar")
def api_yoklamalar():
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    bas = request.args.get("baslangic", date.today().strftime("%Y-%m-%d"))
    bit = request.args.get("bitis",     date.today().strftime("%Y-%m-%d"))
    with get_db() as conn:
        rows = conn.execute("""
            SELECT y.tarih, p.ad_soyad,
                   y.giris_saati, y.cikis_saati,
                   y.giris_uyari, y.cikis_uyari, y.cikis_disari
            FROM yoklama y
            JOIN personel p ON p.id=y.personel_id
            WHERE y.tarih BETWEEN ? AND ?
            ORDER BY y.tarih DESC, p.ad_soyad
        """, (bas, bit)).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/admin/aylik_rapor")
def api_aylik_rapor():
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    yil = int(request.args.get("yil",  date.today().year))
    ay  = int(request.args.get("ay",   date.today().month))
    bas = f"{yil}-{ay:02d}-01"
    bit = f"{yil}-{ay:02d}-31"
    with get_db() as conn:
        rows = conn.execute("""
            SELECT p.id, p.ad_soyad,
                   y.tarih, y.giris_saati, y.cikis_saati,
                   y.giris_uyari, y.cikis_uyari, y.cikis_disari
            FROM yoklama y
            JOIN personel p ON p.id=y.personel_id
            WHERE y.tarih BETWEEN ? AND ? AND p.aktif=1
            ORDER BY p.ad_soyad, y.tarih
        """, (bas, bit)).fetchall()

    # Personel bazında grupla
    personeller = {}
    for r in rows:
        pid = r["id"]
        if pid not in personeller:
            personeller[pid] = {
                "ad_soyad":      r["ad_soyad"],
                "toplam_dk":     0,
                "calisilan_gun": 0,
                "eksik_giris":   0,
                "eksik_cikis":   0,
                "mesai_uyari":   0,
                "klinik_disari": 0,
                "gunler":        []
            }
        p = personeller[pid]
        dk, sure_text = sure_hesapla(r["giris_saati"] or "", r["cikis_saati"] or "", r["tarih"])
        if r["giris_saati"]:
            p["calisilan_gun"] += 1
            p["toplam_dk"]     += dk
        if not r["giris_saati"]:  p["eksik_giris"]   += 1
        if not r["cikis_saati"]:  p["eksik_cikis"]   += 1
        if r["giris_uyari"] or r["cikis_uyari"]: p["mesai_uyari"]   += 1
        if r["cikis_disari"]:     p["klinik_disari"] += 1
        p["gunler"].append({
            "tarih":       r["tarih"],
            "giris":       r["giris_saati"][:5] if r["giris_saati"] else "-",
            "cikis":       r["cikis_saati"][:5] if r["cikis_saati"] else "-",
            "sure":        sure_text,
            "giris_uyari": bool(r["giris_uyari"]),
            "cikis_uyari": bool(r["cikis_uyari"]),
            "cikis_disari":bool(r["cikis_disari"]),
        })

    sonuc = []
    for pid, p in personeller.items():
        toplam_saat = f"{p['toplam_dk']//60}s {p['toplam_dk']%60}dk"
        sonuc.append({**p, "toplam_saat": toplam_saat})
    return jsonify(sonuc)

@app.route("/api/admin/qr_urls")
def api_qr_urls():
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    base = request.host_url.rstrip("/")
    return jsonify({"giris_url": f"{base}/giris", "cikis_url": f"{base}/cikis"})

@app.route("/api/admin/konum_ayar")
def api_konum_get():
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    return jsonify({"lat": KLINIK_LAT, "lng": KLINIK_LNG, "yaricap": KLINIK_YARICAP_M})

@app.route("/api/admin/konum_ayar", methods=["POST"])
def api_konum_set():
    if not admin_kontrol(request.args.get("s","")):
        return jsonify({"hata": "Yetkisiz"}), 401
    global KLINIK_LAT, KLINIK_LNG, KLINIK_YARICAP_M
    d = request.get_json()
    try:
        KLINIK_LAT, KLINIK_LNG = float(d["lat"]), float(d["lng"])
        KLINIK_YARICAP_M = int(d["yaricap"])
        return jsonify({"ok": True, "mesaj": f"Güncellendi. Yarıçap: {KLINIK_YARICAP_M} m"})
    except Exception as e:
        return jsonify({"ok": False, "mesaj": str(e)})

# ── EXPORT ROUTES ─────────────────────────────────────────────────────────────
import io
from flask import send_file
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AYLAR = ["","Ocak","Şubat","Mart","Nisan","Mayıs","Haziran",
         "Temmuz","Ağustos","Eylül","Ekim","Kasım","Aralık"]

def excel_stil(ws):
    """Tüm sütunları otomatik genişlet."""
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

def baslik_stili():
    return {
        "font":      Font(bold=True, color="FFFFFF", size=11),
        "fill":      PatternFill("solid", fgColor="1E3A5F"),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border":    Border(bottom=Side(style="thin", color="FFFFFF"))
    }

def uygula_stil(cell, **kwargs):
    for k, v in kwargs.items():
        setattr(cell, k, v)

def ihlaL_renk(row_cells, giris_u, cikis_u, cikis_d):
    """İhlal varsa satırı turuncu/sarı yap."""
    renk = None
    if cikis_d:  renk = "FFF3E0"   # turuncu tonu
    elif giris_u or cikis_u: renk = "FFFDE7"  # sarı tonu
    if renk:
        for c in row_cells:
            c.fill = PatternFill("solid", fgColor=renk)

# ── 1. AYLIK EXCEL ────────────────────────────────────────────────────────────
@app.route("/export/aylik")
def export_aylik():
    if not admin_kontrol(request.args.get("s","")):
        return "Yetkisiz", 401
    yil = int(request.args.get("yil", date.today().year))
    ay  = int(request.args.get("ay",  date.today().month))
    bas, bit = f"{yil}-{ay:02d}-01", f"{yil}-{ay:02d}-31"

    with get_db() as conn:
        rows = conn.execute("""
            SELECT p.ad_soyad, y.tarih, y.giris_saati, y.cikis_saati,
                   y.giris_uyari, y.cikis_uyari, y.cikis_disari
            FROM yoklama y JOIN personel p ON p.id=y.personel_id
            WHERE y.tarih BETWEEN ? AND ? AND p.aktif=1
            ORDER BY p.ad_soyad, y.tarih
        """, (bas, bit)).fetchall()

    wb = openpyxl.Workbook()

    # ── ÖZET SAYFASI ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Özet"
    ws.row_dimensions[1].height = 30

    basliklar = ["Personel","Çalışılan Gün","Toplam Saat","Eksik Çıkış","Mesai İhlali","Klinik Dışı Çıkış","Durum"]
    st = baslik_stili()
    for i, b in enumerate(basliklar, 1):
        c = ws.cell(row=1, column=i, value=b)
        uygula_stil(c, **st)

    # Personel bazında grupla
    from collections import defaultdict
    ozet = defaultdict(lambda: {"gun":0,"dk":0,"eksik_c":0,"mesai":0,"disari":0})
    for r in rows:
        ad = r["ad_soyad"]
        if r["giris_saati"]:
            ozet[ad]["gun"] += 1
            dk, _ = sure_hesapla(r["giris_saati"], r["cikis_saati"] or r["giris_saati"], r["tarih"])
            ozet[ad]["dk"] += dk
        if not r["cikis_saati"] and r["giris_saati"]: ozet[ad]["eksik_c"] += 1
        if r["giris_uyari"] or r["cikis_uyari"]:      ozet[ad]["mesai"]   += 1
        if r["cikis_disari"]:                          ozet[ad]["disari"]  += 1

    for i, (ad, d) in enumerate(ozet.items(), 2):
        saat_txt = f"{d['dk']//60}s {d['dk']%60}dk"
        durum = "✅ Normal" if d["mesai"]==0 and d["disari"]==0 else "⚠️ İhlal Var"
        satirlar = [ad, d["gun"], saat_txt, d["eksik_c"], d["mesai"], d["disari"], durum]
        for j, v in enumerate(satirlar, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(horizontal="center")
            if d["disari"] > 0 or d["mesai"] > 0:
                c.fill = PatternFill("solid", fgColor="FFF3E0")
    excel_stil(ws)

    # ── DETAY SAYFASI ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Günlük Detay")
    basliklar2 = ["Personel","Tarih","Gün","Giriş","Çıkış","Çalışma Süresi","Giriş İhlali","Çıkış İhlali","Klinik Dışı"]
    for i, b in enumerate(basliklar2, 1):
        c = ws2.cell(row=1, column=i, value=b)
        uygula_stil(c, **st)

    gun_adlari = ["Pazartesi","Salı","Çarşamba","Perşembe","Cuma","Cumartesi","Pazar"]
    for ri, r in enumerate(rows, 2):
        dk, sure_txt = sure_hesapla(r["giris_saati"] or "", r["cikis_saati"] or "", r["tarih"])
        try: gun = gun_adlari[date.fromisoformat(r["tarih"]).weekday()]
        except: gun = ""
        vals = [
            r["ad_soyad"], r["tarih"], gun,
            r["giris_saati"][:5] if r["giris_saati"] else "-",
            r["cikis_saati"][:5] if r["cikis_saati"] else "-",
            sure_txt,
            "Evet" if r["giris_uyari"] else "Hayır",
            "Evet" if r["cikis_uyari"] else "Hayır",
            "Evet" if r["cikis_disari"] else "Hayır",
        ]
        row_cells = [ws2.cell(row=ri, column=j, value=v) for j, v in enumerate(vals, 1)]
        for c in row_cells: c.alignment = Alignment(horizontal="center")
        ihlaL_renk(row_cells, r["giris_uyari"], r["cikis_uyari"], r["cikis_disari"])
    excel_stil(ws2)

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    dosya_adi = f"yoklama_{yil}_{ay:02d}_{AYLAR[ay]}.xlsx"
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=dosya_adi)

# ── 2. TARİH ARALIKLI CSV ─────────────────────────────────────────────────────
@app.route("/export/csv")
def export_csv():
    if not admin_kontrol(request.args.get("s","")):
        return "Yetkisiz", 401
    bas = request.args.get("bas", date.today().strftime("%Y-%m-%d"))
    bit = request.args.get("bit", date.today().strftime("%Y-%m-%d"))

    with get_db() as conn:
        rows = conn.execute("""
            SELECT p.ad_soyad, y.tarih, y.giris_saati, y.cikis_saati,
                   y.giris_uyari, y.cikis_uyari, y.cikis_disari
            FROM yoklama y JOIN personel p ON p.id=y.personel_id
            WHERE y.tarih BETWEEN ? AND ?
            ORDER BY y.tarih DESC, p.ad_soyad
        """, (bas, bit)).fetchall()

    import csv
    buf = io.StringIO()
    yaz = csv.writer(buf)
    yaz.writerow(["Personel","Tarih","Giriş","Çıkış","Çalışma Süresi","Giriş İhlali","Çıkış İhlali","Klinik Dışı"])
    for r in rows:
        _, sure = sure_hesapla(r["giris_saati"] or "", r["cikis_saati"] or "", r["tarih"])
        yaz.writerow([
            r["ad_soyad"], r["tarih"],
            r["giris_saati"][:5] if r["giris_saati"] else "-",
            r["cikis_saati"][:5] if r["cikis_saati"] else "-",
            sure,
            "Evet" if r["giris_uyari"] else "Hayır",
            "Evet" if r["cikis_uyari"] else "Hayır",
            "Evet" if r["cikis_disari"] else "Hayır",
        ])

    buf.seek(0)
    bbuf = io.BytesIO(buf.getvalue().encode("utf-8-sig"))  # utf-8-sig = Excel'de Türkçe karakter sorunu olmaz
    dosya_adi = f"yoklama_{bas}_{bit}.csv"
    return send_file(bbuf, mimetype="text/csv", as_attachment=True, download_name=dosya_adi)

# ── 3. TÜM ZAMANLAR EXCEL ─────────────────────────────────────────────────────
@app.route("/export/tumü")
def export_tumu():
    if not admin_kontrol(request.args.get("s","")):
        return "Yetkisiz", 401

    with get_db() as conn:
        rows = conn.execute("""
            SELECT p.ad_soyad, y.tarih, y.giris_saati, y.cikis_saati,
                   y.giris_uyari, y.cikis_uyari, y.cikis_disari
            FROM yoklama y JOIN personel p ON p.id=y.personel_id
            ORDER BY y.tarih DESC, p.ad_soyad
        """).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tüm Kayıtlar"
    st = baslik_stili()
    basliklar = ["Personel","Tarih","Gün","Giriş","Çıkış","Çalışma Süresi","Giriş İhlali","Çıkış İhlali","Klinik Dışı"]
    for i, b in enumerate(basliklar, 1):
        c = ws.cell(row=1, column=i, value=b)
        uygula_stil(c, **st)

    gun_adlari = ["Pazartesi","Salı","Çarşamba","Perşembe","Cuma","Cumartesi","Pazar"]
    for ri, r in enumerate(rows, 2):
        _, sure = sure_hesapla(r["giris_saati"] or "", r["cikis_saati"] or "", r["tarih"])
        try: gun = gun_adlari[date.fromisoformat(r["tarih"]).weekday()]
        except: gun = ""
        vals = [
            r["ad_soyad"], r["tarih"], gun,
            r["giris_saati"][:5] if r["giris_saati"] else "-",
            r["cikis_saati"][:5] if r["cikis_saati"] else "-",
            sure,
            "Evet" if r["giris_uyari"] else "Hayır",
            "Evet" if r["cikis_uyari"] else "Hayır",
            "Evet" if r["cikis_disari"] else "Hayır",
        ]
        row_cells = [ws.cell(row=ri, column=j, value=v) for j, v in enumerate(vals, 1)]
        for c in row_cells: c.alignment = Alignment(horizontal="center")
        ihlaL_renk(row_cells, r["giris_uyari"], r["cikis_uyari"], r["cikis_disari"])
    excel_stil(ws)

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    dosya_adi = f"yoklama_tum_kayitlar_{date.today()}.xlsx"
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=dosya_adi)

# ── MAIN ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    init_db()
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
